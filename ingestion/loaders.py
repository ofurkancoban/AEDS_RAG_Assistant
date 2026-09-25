import csv
from pathlib import Path

from langchain_core.documents import Document


MIN_TABLE_ROWS = 3  # header + at least 2 data rows

MAX_FILTERABLE_FIELD_LENGTH = 100


def _row_field_metadata(pairs: list[tuple[str, str]]) -> dict[str, str]:
    """Expose short, scalar column values (e.g. 'compulsory: yes', 'language:
    English') as their own 'field_<column>' metadata entries, so a structured
    row's own attributes can be queried directly via an exact metadata filter
    (e.g. field_compulsory == "yes") - not just found by similarity ranking,
    which struggles to reliably surface every matching row out of many when a
    question asks to enumerate all of them rather than find the most relevant
    few. Long free-text columns (descriptions, skills) are left out since
    they're not meaningful as an exact-match filter."""
    metadata = {}
    for key, value in pairs:
        value = value.strip()
        if not value or len(value) > MAX_FILTERABLE_FIELD_LENGTH:
            continue
        safe_key = "".join(c if c.isalnum() else "_" for c in key.strip().lower())
        metadata[f"field_{safe_key}"] = value
    return metadata


def _table_to_markdown(table: list[list[str | None]]) -> str:
    rows = [[(cell or "").strip().replace("\n", " ") for cell in row] for row in table]
    rows = [row for row in rows if any(cell for cell in row)]
    if len(rows) < MIN_TABLE_ROWS:
        # Below this, extractions are usually stray labelled boxes or graphic
        # design elements pdfplumber misreads as a table, not real tabular data.
        return ""

    header, *body = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def load_pdf(path: Path) -> list[Document]:
    """Extract page text with pypdf (fast, reliable for prose) and, separately,
    any detected tables with pdfplumber rendered as markdown tables. Tables are
    kept as their own documents (flagged is_table=True) so the chunker treats
    them as atomic units instead of splitting them mid-row, which otherwise
    scrambles which figure belongs to which row/column in table-heavy PDFs
    (e.g. a module handbook) once flattened to plain text."""
    import pdfplumber
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    documents = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={"source_file": path.name, "page": page_number},
                )
            )

    with pdfplumber.open(str(path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                markdown_table = _table_to_markdown(table)
                if markdown_table:
                    documents.append(
                        Document(
                            page_content=markdown_table,
                            metadata={"source_file": path.name, "page": page_number, "is_table": True},
                        )
                    )

    return documents


def load_markdown_or_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8")
    return [Document(page_content=text, metadata={"source_file": path.name})]


def load_docx(path: Path) -> list[Document]:
    import docx

    doc = docx.Document(str(path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    return [Document(page_content=text, metadata={"source_file": path.name})]


def load_xlsx(path: Path) -> list[Document]:
    """One Document per row, formatted as labelled 'column: value' pairs using
    the header row. A plain '|'-joined row loses which value belongs to which
    column once it's just one line among many in a chunk with no header in
    sight - labelling every row keeps each one self-contained and meaningful
    on its own, which matters a lot once rows are split across chunks."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True)
    documents = []
    for sheet in workbook.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            continue
        header = [str(cell) if cell is not None else f"column_{i}" for i, cell in enumerate(header)]

        for row_number, row in enumerate(rows_iter, start=2):
            row_pairs = [(key, str(value)) for key, value in zip(header, row) if value is not None]
            if row_pairs:
                documents.append(
                    Document(
                        page_content="\n".join(f"{key}: {value}" for key, value in row_pairs),
                        metadata={
                            "source_file": path.name,
                            "sheet": sheet.title,
                            "row": row_number,
                            **_row_field_metadata(row_pairs),
                        },
                    )
                )
    return documents


def load_pptx(path: Path) -> list[Document]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    documents = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if not texts:
            continue
        documents.append(
            Document(
                page_content="\n".join(texts),
                metadata={"source_file": path.name, "slide": slide_number},
            )
        )
    return documents


def load_html(path: Path) -> list[Document]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    return [Document(page_content=text, metadata={"source_file": path.name})]


def load_csv(path: Path) -> list[Document]:
    """One Document per row, as labelled 'column: value' pairs (see load_xlsx
    docstring for why this matters more than it looks)."""
    documents = []
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row_number, row in enumerate(reader, start=1):
            row_pairs = [(key, value) for key, value in row.items() if value]
            if row_pairs:
                documents.append(
                    Document(
                        page_content="\n".join(f"{key}: {value}" for key, value in row_pairs),
                        metadata={
                            "source_file": path.name,
                            "row": row_number,
                            **_row_field_metadata(row_pairs),
                        },
                    )
                )
    return documents


LOADERS = {
    ".pdf": load_pdf,
    ".md": load_markdown_or_text,
    ".txt": load_markdown_or_text,
    ".docx": load_docx,
    ".xlsx": load_xlsx,
    ".pptx": load_pptx,
    ".html": load_html,
    ".htm": load_html,
    ".csv": load_csv,
}


def load_document(path: Path) -> list[Document]:
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return loader(path)
